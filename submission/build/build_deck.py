#!/usr/bin/env python3
"""Build the HOPTRACE project presentation (PPTX).

    python submission/build/build_deck.py [--slot F1 --team 7]

Native PowerPoint shapes, connectors, tables and the genuine screenshots in
submission/screenshots/. Visual identity follows the HOPTRACE website
(pale dotted background, dark outlines, rounded cards, offset shadows,
cyan / coral / yellow / lime / violet accents).

All headline numbers are asserted against evaluate/*.json before the deck is written.
Documentation-only dependencies: python-pptx, Pillow.
"""
from __future__ import annotations

import argparse
import copy
import json
import os

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.abspath(os.path.join(HERE, "..", ".."))
SHOTS = os.path.join(ROOT, "submission", "screenshots")
FIGS = os.path.join(ROOT, "submission", "figures")

# ----------------------------------------------------------------- palette / fonts
INK, SOFT = "14181A", "3D4D55"
MINT50, MINT100, MINT200 = "F3F8F6", "E7EFF2", "D3E3E6"
SIGNAL, STRUCT, HEALTHY, DEGRADED, UNKNOWN = "2FD3E0", "6A4DFF", "B4FF3D", "FF6B4A", "FFD23F"
WHITE = "FFFFFF"
LILAC, CYAN_T, CORAL_T, YEL_T = "E5DEFE", "D9F7FA", "FFE3DC", "FFF4C4"
HEAD, BODY, MONO = "Arial Black", "Calibri", "Consolas"

SW, SH = 13.333, 7.5


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


# ----------------------------------------------------------------- numbers
def check_numbers() -> None:
    a = json.load(open(os.path.join(ROOT, "evaluate", "held_out_results.json")))["summary"]
    b = json.load(open(os.path.join(ROOT, "evaluate", "held_out_results_4b.json")))["summary"]
    f = lambda s, k: (s[k]["numerator"], s[k]["denominator"])
    assert f(a, "false_attribution_rate") == (3, 12) and f(b, "false_attribution_rate") == (1, 12)
    assert f(a, "clean_control_false_positive_rate") == (2, 4) and f(b, "clean_control_false_positive_rate") == (0, 4)
    assert f(a, "within_1_hop_accuracy") == (3, 3) and f(b, "within_1_hop_accuracy") == (4, 4)
    assert f(a, "exact_hop_accuracy") == (2, 3) and f(b, "exact_hop_accuracy") == (3, 4)
    assert f(a, "detection_recall") == (3, 8) and f(b, "detection_recall") == (4, 8)
    ba = json.load(open(os.path.join(ROOT, "evaluate", "held_out_results.json")))["brier_score"]
    bb = json.load(open(os.path.join(ROOT, "evaluate", "held_out_results_4b.json")))["brier_score"]
    assert (ba["n"], round(ba["brier"], 3)) == (5, 0.432) and (bb["n"], round(bb["brier"], 3)) == (4, 0.04)
    print("deck numbers match evaluate/*.json")


# ----------------------------------------------------------------- low-level helpers
class Deck:
    def __init__(self, slot: str, team: str):
        self.prs = Presentation()
        self.prs.slide_width = Inches(SW)
        self.prs.slide_height = Inches(SH)
        self.slot, self.team = slot, team
        self._set_master_background()
        self.n = 0
        self.total = 18

    def _set_master_background(self):
        master = self.prs.slide_master
        _, rid = master.part.get_or_add_image_part(os.path.join(HERE, "bg_dots.png"))
        ns = {"p": "http://schemas.openxmlformats.org/presentationml/2006/main",
              "a": "http://schemas.openxmlformats.org/drawingml/2006/main",
              "r": "http://schemas.openxmlformats.org/officeDocument/2006/relationships"}
        xml = (f'<p:bg xmlns:p="{ns["p"]}" xmlns:a="{ns["a"]}" xmlns:r="{ns["r"]}"><p:bgPr>'
               f'<a:blipFill dpi="0" rotWithShape="1"><a:blip r:embed="{rid}"/><a:srcRect/>'
               f'<a:stretch><a:fillRect/></a:stretch></a:blipFill><a:effectLst/></p:bgPr></p:bg>')
        cSld = master.element.find("{%s}cSld" % ns["p"])
        old = cSld.find("{%s}bg" % ns["p"])
        if old is not None:
            cSld.remove(old)
        cSld.insert(0, etree.fromstring(xml))

    # -- slide creation -------------------------------------------------
    def slide(self, title: str | None, notes: str, footer: bool = True):
        layout = self.prs.slide_layouts[5]  # Title Only
        s = self.prs.slides.add_slide(layout)
        self.n += 1
        t = s.shapes.title
        if title is None:
            t._element.getparent().remove(t._element)
        else:
            t.left, t.top, t.width, t.height = Inches(0.6), Inches(0.32), Inches(12.1), Inches(0.95)
            tf = t.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = tf.margin_right = Inches(0.0)
            p = tf.paragraphs[0]
            p.alignment = PP_ALIGN.LEFT
            r = p.add_run()
            r.text = title
            r.font.name, r.font.size, r.font.bold = HEAD, Pt(32), False
            r.font.color.rgb = rgb(INK)
        s.notes_slide.notes_text_frame.text = notes
        if footer:
            self.text(s, 0.6, 7.06, 9.0, 0.3, ["HOPTRACE · BCSE308P Computer Networks Lab · Krish Chirag Shah · 24BCE0624"],
                      size=10, color=SOFT)
            self.text(s, 11.7, 7.06, 1.03, 0.3, [f"{self.n}"], size=11, color=INK, bold=True, align="r", mono=True)
        return s

    # -- primitives -----------------------------------------------------
    def _fmt_run(self, r, size, color, bold, italic, mono, font=None):
        r.font.name = font or (MONO if mono else BODY)
        r.font.size = Pt(size)
        r.font.bold = bold
        r.font.italic = italic
        r.font.color.rgb = rgb(color)

    def text(self, s, x, y, w, h, paras, size=16, color=INK, bold=False, italic=False, align="l", mono=False,
             anchor="t", margin=0.04, line=None, space_after=0, font=None, name=None):
        """paras: list of paragraphs; a paragraph is a str or a list of (text, {opts}) runs."""
        tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        if name:
            tb.name = name
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = Inches(margin)
        tf.margin_top = tf.margin_bottom = Inches(0.02)
        tf.vertical_anchor = {"t": MSO_ANCHOR.TOP, "m": MSO_ANCHOR.MIDDLE, "b": MSO_ANCHOR.BOTTOM}[anchor]
        for i, para in enumerate(paras):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            if line:
                p.line_spacing = line
            p.space_after = Pt(space_after)
            runs = [(para, {})] if isinstance(para, str) else para
            for txt, o in runs:
                r = p.add_run()
                r.text = txt
                self._fmt_run(r, o.get("size", size), o.get("color", color), o.get("bold", bold),
                              o.get("italic", italic), o.get("mono", mono), o.get("font", font))
                if o.get("link"):
                    r.hyperlink.address = o["link"]
        return tb

    def card(self, s, x, y, w, h, fill=WHITE, outline=INK, shadow=True, dash=False, radius=0.12, lw=2.25, name=None):
        if shadow:
            sh = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x + 0.07), Inches(y + 0.07), Inches(w), Inches(h))
            sh.adjustments[0] = radius
            sh.fill.solid(); sh.fill.fore_color.rgb = rgb(INK)
            sh.line.fill.background()
            sh.shadow.inherit = False
            sh.name = (name or "card") + " shadow"
        c = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        c.adjustments[0] = radius
        c.fill.solid(); c.fill.fore_color.rgb = rgb(fill)
        c.line.color.rgb = rgb(outline); c.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            c.line.dash_style = MSO_LINE.DASH
        c.shadow.inherit = False
        c.name = name or "card"
        return c

    def pill(self, s, x, y, w, h, label, fill=WHITE, size=12, bold=True, mono=False, dash=False, color=INK, shadow=False):
        c = self.card(s, x, y, w, h, fill=fill, shadow=shadow, dash=dash, radius=0.5, lw=1.75, name="pill")
        tf = c.text_frame
        tf.margin_left = tf.margin_right = Inches(0.06); tf.margin_top = tf.margin_bottom = Inches(0.01)
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        tf.word_wrap = True
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        self._fmt_run(r, size, color, bold, False, mono)
        return c

    def circle(self, s, cx, cy, d, label, fill, size=16, color=INK, outline=INK, dash=False, lw=2.25):
        o = s.shapes.add_shape(MSO_SHAPE.OVAL, Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
        o.fill.solid(); o.fill.fore_color.rgb = rgb(fill)
        o.line.color.rgb = rgb(outline); o.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            o.line.dash_style = MSO_LINE.DASH
        o.shadow.inherit = False
        tf = o.text_frame
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
        r = p.add_run(); r.text = label
        self._fmt_run(r, size, color, True, False, False, HEAD)
        return o

    def line(self, s, x1, y1, x2, y2, color=INK, width=2.5, dash=False, arrow=False):
        c = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(width)
        if dash:
            from pptx.enum.dml import MSO_LINE
            c.line.dash_style = MSO_LINE.DASH
        if arrow:
            ln = c.line._get_or_add_ln()
            te = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd")
            te.set("type", "triangle"); te.set("w", "med"); te.set("len", "med")
        c.name = "arrow" if arrow else "line"
        return c

    def picture(self, s, name, x, y, w=None, h=None, alt="", border=True, folder=SHOTS):
        path = os.path.join(folder, name)
        kw = {}
        if w:
            kw["width"] = Inches(w)
        if h:
            kw["height"] = Inches(h)
        pic = s.shapes.add_picture(path, Inches(x), Inches(y), **kw)
        pic._element.nvPicPr.cNvPr.set("descr", alt)
        pic.name = alt[:40] or name
        if border:
            pic.line.color.rgb = rgb(INK); pic.line.width = Pt(2)
        return pic

    def code(self, s, x, y, w, h, lines, size=12):
        c = self.card(s, x, y, w, h, fill=INK, outline=INK, shadow=False, radius=0.08, lw=1)
        self.text(s, x + 0.12, y + 0.08, w - 0.24, h - 0.16, lines, size=size, color="E8F1F3", mono=True, anchor="m", margin=0)
        return c

    def sticker(self, s, x, y, w, label, fill=UNKNOWN, size=12):
        return self.pill(s, x, y, w, 0.34, label, fill=fill, size=size, mono=True, shadow=False)

    def table(self, s, x, y, w, col_w, rows, header_fill=INK, size=14, row_h=0.42, fills=None, bold_first_col=False, mono_cols=()):
        nrows, ncols = len(rows), len(rows[0])
        gs = s.shapes.add_table(nrows, ncols, Inches(x), Inches(y), Inches(w), Inches(row_h * nrows))
        tbl = gs.table
        tblPr = tbl._tbl.tblPr
        for attr in ("bandRow", "firstRow"):
            tblPr.set(attr, "0")
        sid = tblPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}tableStyleId")
        if sid is not None:
            sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"  # "No Style, No Grid"
        for i, cw in enumerate(col_w):
            tbl.columns[i].width = Inches(cw)
        for r in range(nrows):
            tbl.rows[r].height = Inches(row_h)
            for c in range(ncols):
                cell = tbl.cell(r, c)
                cell.margin_left = cell.margin_right = Inches(0.08)
                cell.margin_top = cell.margin_bottom = Inches(0.03)
                cell.vertical_anchor = MSO_ANCHOR.MIDDLE
                val = rows[r][c]
                tf = cell.text_frame
                tf.word_wrap = True
                p = tf.paragraphs[0]
                p.alignment = PP_ALIGN.LEFT if c == 0 else PP_ALIGN.CENTER
                run = p.add_run(); run.text = val
                head = r == 0
                fill = header_fill if head else (fills[r][c] if fills and fills[r][c] else (WHITE if r % 2 else MINT50))
                cell.fill.solid(); cell.fill.fore_color.rgb = rgb(fill)
                self._fmt_run(run, size, WHITE if head else INK, head or (bold_first_col and c == 0), False,
                              (c in mono_cols and not head))
                # thin borders
                tcPr = cell._tc.get_or_add_tcPr()
                for tag in ("lnL", "lnR", "lnT", "lnB"):
                    ln = etree.SubElement(tcPr, "{http://schemas.openxmlformats.org/drawingml/2006/main}%s" % tag)
                    ln.set("w", "12700")
                    sf = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
                    clr = etree.SubElement(sf, "{http://schemas.openxmlformats.org/drawingml/2006/main}srgbClr")
                    clr.set("val", "9DB0B7" if not head else INK)
                # borders must precede the fill in tcPr
                fill_el = tcPr.find("{http://schemas.openxmlformats.org/drawingml/2006/main}solidFill")
                if fill_el is not None:
                    tcPr.remove(fill_el); tcPr.append(fill_el)
        return gs


# ----------------------------------------------------------------- slides
def build(slot: str, team: str) -> str:
    check_numbers()
    D = Deck(slot, team)
    S = D  # alias

    # ============================================================ 1 TITLE
    s = D.slide(None, "Introduce the project and the central claim: HOPTRACE estimates where along an observed network path an "
                      "application's degradation appears to begin. Mention up front that it is a research prototype and that the "
                      "current detector has not been independently validated.", footer=False)
    D.n = 1
    D.pill(s, 0.7, 0.6, 7.5, 0.4, "BCSE308P · Computer Networks Lab · Innovative Evaluation Practices", size=13, mono=True, shadow=True)
    D.text(s, 0.65, 1.25, 8.2, 1.6, ["HOPTRACE"], size=74, font=HEAD, color=INK, anchor="m")
    D.text(s, 0.7, 2.95, 7.4, 1.2, ["Network Path Fault Localization through Application-to-Hop Correlation"], size=26, bold=True, color=INK)
    # student card
    D.card(s, 0.7, 4.35, 7.2, 1.75, fill=WHITE, name="student card")
    D.text(s, 0.95, 4.47, 6.8, 1.55, [
        [("Krish Chirag Shah", {"bold": True, "size": 22}), ("   sole author", {"size": 14, "color": SOFT, "italic": True})],
        [("Reg. no. 24BCE0624", {"mono": True, "size": 16})],
        [("Assignment 9 · Project Presentation (VTOP)", {"size": 16})],
        [("Slot ", {"size": 16}), (slot, {"bold": True, "size": 16, "mono": True}),
         ("   ·   Team ", {"size": 16}), (team, {"bold": True, "size": 16, "mono": True})],
    ], size=16, space_after=2)
    if slot.startswith("TEMP") or team.startswith("TEMP"):
        D.pill(s, 5.55, 5.55, 2.2, 0.36, "TO BE CONFIRMED", fill=UNKNOWN, size=12, mono=True, dash=True)
    # motif: path of hops (right)
    xs = [9.0, 10.15, 11.3, 12.45]
    D.text(s, 8.0, 1.0, 5.0, 0.4, ["a path, seen hop by hop"], size=14, color=SOFT, align="c", italic=True)
    cy = 2.4
    cols = [HEALTHY, HEALTHY, DEGRADED, DEGRADED, DEGRADED]
    pts = [(8.95, 1.95), (9.9, 2.4), (10.85, 2.85), (11.8, 3.3), (12.4, 3.85)]
    D.circle(s, 8.15, 1.55, 0.7, "C", STRUCT, size=16, color=WHITE)
    D.line(s, 8.15 + 0.3, 1.65, pts[0][0] - 0.3, pts[0][1] - 0.05, width=4)
    for i in range(1, 5):
        D.line(s, pts[i - 1][0], pts[i - 1][1], pts[i][0], pts[i][1], width=4, dash=(i >= 3))
    for i, (px, py) in enumerate(pts):
        D.circle(s, px, py, 0.62, str(i + 1), cols[i], size=18)
    D.pill(s, 9.4, 4.75, 3.6, 0.5, "appears to begin at or after Hop 3", fill=SIGNAL, size=14, shadow=True)
    D.text(s, 8.7, 5.4, 4.3, 0.5, ["Illustration of the idea; not measured data."], size=12, color=SOFT, align="c", italic=True)
    # banner
    D.card(s, 0.7, 6.4, 12.0, 0.7, fill=UNKNOWN, dash=True, shadow=False, name="status banner")
    D.text(s, 0.9, 6.43, 11.6, 0.65, [[("Research prototype. ", {"bold": True}),
                                         ("The current detector has not been independently validated; the independent Phase 4 run missed its "
                                          "false-attribution target (3/12 = 25% vs at most 10%).", {})]], size=14, anchor="m")

    # ============================================================ 2 PROBLEM
    s = D.slide("Slow, but where on the path?",
                "The problem: when an application slows down, ping and traceroute describe the path but not what the application "
                "experienced, and application monitoring sees the symptom but not the path. HOPTRACE puts both on one clock. "
                "Stress the boundary of the claim: when changes appear, not why.")
    cw, gap, y0, ch = 3.85, 0.3, 1.5, 3.4
    heads = [("ping and traceroute", "measure the path", ["Per-hop reachability and RTT", "Nothing about what the application felt"], WHITE),
             ("Application monitoring", "measures the symptom", ["Sees that responses got slower", "Nothing about where on the path"], WHITE),
             ("HOPTRACE", "joins the two on one clock", ["Probes every hop and samples the app's own RTT", "Compares when each series starts to misbehave"], CYAN_T)]
    for i, (h1, h2, bl, fill) in enumerate(heads):
        x = 0.6 + i * (cw + gap)
        D.card(s, x, y0, cw, ch, fill=fill)
        D.text(s, x + 0.25, y0 + 0.2, cw - 0.5, 0.5, [h1], size=24, bold=True)
        D.text(s, x + 0.25, y0 + 0.78, cw - 0.5, 0.4, [h2], size=17, color=SOFT, italic=True)
        D.text(s, x + 0.25, y0 + 1.5, cw - 0.5, 1.85, [[("•  ", {"bold": True}), (b, {})] for b in bl], size=19, space_after=10)
    D.card(s, 0.6, 5.3, 12.1, 1.55, fill=WHITE, shadow=True, name="question")
    D.text(s, 0.85, 5.36, 11.6, 1.43, [
        [("The question: ", {"bold": True}), ("where along the observed path does the application's degradation first appear?", {})],
        [("Not the question: ", {"bold": True, "color": DEGRADED}), ("which router caused it. Timing correlation does not establish causation.", {})]],
        size=21, anchor="m", space_after=6)

    # ============================================================ 3 OBJECTIVES
    s = D.slide("Objectives and scope",
                "Five objectives on the left. On the right, what is in scope and, just as importantly, what is explicitly out of scope "
                "and not claimed: real networks, calibrated confidence, and reliable localization of loss and bandwidth faults.")
    objs = ["Build a reproducible testbed with a known path and injectable faults",
            "Measure path hops and application RTT on one shared clock",
            "Detect onsets, align them in time, report evidence-bounded results",
            "Keep ground truth out of the engine's reach; score afterwards",
            "Evaluate on separate tuning and held-out sets, and disclose misses"]
    for i, o in enumerate(objs):
        y = 1.55 + i * 1.05
        D.circle(s, 0.95, y + 0.4, 0.62, str(i + 1), [SIGNAL, UNKNOWN, HEALTHY, DEGRADED, STRUCT][i], size=18,
                 color=WHITE if i == 4 else INK)
        D.text(s, 1.45, y, 5.4, 0.85, [o], size=17, anchor="m")
    D.card(s, 7.2, 1.5, 5.5, 2.65, fill=CYAN_T, name="in scope")
    D.text(s, 7.45, 1.58, 5.0, 0.4, ["In scope (implemented)"], size=19, bold=True)
    D.text(s, 7.45, 2.05, 5.0, 2.1, [[("•  ", {"bold": True}), (t, {})] for t in
           ["6-namespace linear testbed, 5 hops", "Raw-socket probing + WebSocket telemetry",
            "Latency, jitter, loss, bandwidth faults", "Median/MAD onsets, alignment, typed evidence",
            "Stored-result dashboard, guided explanation"]], size=16, space_after=4)
    D.card(s, 7.2, 4.45, 5.5, 2.4, fill=CORAL_T, dash=True, shadow=False, name="out of scope")
    D.text(s, 7.45, 4.52, 5.0, 0.4, ["Out of scope (not claimed)"], size=19, bold=True)
    D.text(s, 7.45, 4.98, 5.0, 1.85, [[("•  ", {"bold": True}), (t, {})] for t in
           ["Real or multi-path networks", "Reliable loss / bandwidth localization",
            "Calibrated confidence; causal claims", "Sealed, full-scale evaluation (spec Phase 5)"]], size=16, space_after=4)

    # ============================================================ 4 SOLUTION
    s = D.slide("Key idea: measure, detect, align, report",
                "The four-step loop. Measure on one monotonic clock; detect an onset per series using median plus five MAD, requiring three "
                "buckets in a row; align the application onset with hop onsets inside a 1.5 second window, requiring downstream persistence; "
                "report using a locked, evidence-bounded sentence.")
    steps = [("1", "Measure", "Probe every hop; sample the app's own RTT", "one CLOCK_MONOTONIC", SIGNAL),
             ("2", "Detect", "Find when each series starts to misbehave", "median + 5·MAD, 3 in a row", UNKNOWN),
             ("3", "Align", "Does a hop's onset line up with the app's?", "|Δt| ≤ 1.5 s + persistence", HEALTHY),
             ("4", "Report", "Name the earliest corroborated hop", "at or after Hop N", LILAC)]
    cw, gap, y0 = 2.65, 0.5, 1.6
    for i, (n, h1, desc, mono, fill) in enumerate(steps):
        x = 0.6 + i * (cw + gap)
        D.card(s, x, y0, cw, 3.4, fill=WHITE)
        D.circle(s, x + 0.55, y0 + 0.62, 0.7, n, fill, size=22)
        D.text(s, x + 1.0, y0 + 0.35, cw - 1.1, 0.6, [h1], size=24, bold=True, anchor="m")
        D.text(s, x + 0.25, y0 + 1.3, cw - 0.5, 1.2, [desc], size=18)
        D.pill(s, x + 0.2, y0 + 2.45, cw - 0.4, 0.75, mono, fill=MINT100, size=14, mono=True, bold=False)
        if i < 3:
            D.line(s, x + cw + 0.04, y0 + 1.7, x + cw + gap - 0.04, y0 + 1.7, width=3.5, arrow=True)
    D.card(s, 0.6, 5.3, 12.1, 1.55, fill=INK, outline=INK, shadow=False, name="output band")
    D.text(s, 0.9, 5.37, 11.5, 0.5, ["The only sentence it will say"], size=14, color=SIGNAL, mono=True)
    D.text(s, 0.9, 5.82, 11.5, 0.6, ["“Degradation appears to begin at or after Hop N.”"], size=28, color=WHITE, bold=True)
    D.text(s, 0.9, 6.42, 11.5, 0.4, ["Never: “Router N caused it.”  Forbidden phrases are asserted out in code."], size=15, color=MINT200)

    # ============================================================ 5 ARCHITECTURE
    s = D.slide("System architecture",
                "Two halves kept apart on purpose. Top: observation and analysis, which reads no ground truth. Bottom: experiment control and "
                "evaluation. The fault controller injects real impairments into the testbed and records labels in a permission-isolated "
                "store; the evaluation joins results to labels only afterwards. The dashboard shows stored bundles only.")
    D.text(s, 0.6, 1.3, 8.0, 0.3, ["OBSERVATION AND ANALYSIS (reads no ground truth)"], size=12, color=SOFT, mono=True, bold=True)
    cw, ch, gap = 2.0, 1.1, 0.55
    xs = [0.6 + i * (cw + gap) for i in range(5)]
    def node(x, y, h, title, sub, fill=WHITE, subsize=12):
        D.card(s, x, y, cw, h, fill=fill)
        D.text(s, x + 0.08, y + 0.08, cw - 0.16, h - 0.16, [[(title, {"bold": True, "size": 15})], [(sub, {"mono": True, "size": subsize, "color": SOFT})]],
               size=15, align="c", anchor="m")
    node(xs[0], 2.3, ch, "Controlled testbed", "6 namespaces", CYAN_T, 12)
    node(xs[1], 1.65, 0.95, "Path prober", "→ hop.jsonl")
    node(xs[1], 2.95, 0.95, "App telemetry", "→ app.jsonl")
    node(xs[2], 2.3, ch, "Correlation engine", "→ EvidenceRecord")
    node(xs[3], 2.3, ch, "Interpretation", "→ path_bundle.json", subsize=11)
    node(xs[4], 2.3, ch, "Dashboard", "reads stored bundles", LILAC, 11)
    ymid = 2.3 + ch / 2
    D.line(s, xs[0] + cw, ymid - 0.2, xs[1], 2.12, arrow=True, width=3)
    D.line(s, xs[0] + cw, ymid + 0.2, xs[1], 3.42, arrow=True, width=3)
    D.line(s, xs[1] + cw, 2.12, xs[2], ymid - 0.2, arrow=True, width=3)
    D.line(s, xs[1] + cw, 3.42, xs[2], ymid + 0.2, arrow=True, width=3)
    D.line(s, xs[2] + cw, ymid, xs[3], ymid, arrow=True, width=3)
    D.line(s, xs[3] + cw, ymid, xs[4], ymid, arrow=True, width=3)
    # boundary
    D.line(s, 0.45, 4.4, 12.9, 4.4, color=DEGRADED, width=2.5, dash=True)
    D.pill(s, 2.35, 4.22, 6.5, 0.36, "✕  isolation boundary: engine and dashboard never read ground truth", fill=WHITE, size=12.5, dash=False)
    D.text(s, 3.6, 4.68, 5.5, 0.3, ["EXPERIMENT CONTROL AND EVALUATION"], size=12, color=SOFT, mono=True, bold=True)
    node(xs[0], 5.35, ch, "Fault controller", "tc netem / tbf", CORAL_T)
    node(xs[2], 5.35, ch, "Ground-truth DB", "OS-permission isolated", YEL_T, 11)
    node(xs[3], 5.35, ch, "Evaluation", "scores vs labels")
    D.line(s, xs[0] + cw / 2, 5.35, xs[0] + cw / 2, 2.3 + ch + 0.03, color=DEGRADED, width=3, dash=True, arrow=True)
    D.pill(s, xs[0] + cw / 2 + 0.12, 4.62, 1.5, 0.3, "injects faults", size=12, shadow=False)
    D.line(s, xs[0] + cw, 5.9, xs[2], 5.9, arrow=True, width=3)
    D.text(s, xs[0] + cw + 0.02, 5.48, gap + cw - 0.05, 0.4, ["writes labels"], size=12, align="c", bold=True)
    D.line(s, xs[2] + cw, 5.9, xs[3], 5.9, arrow=True, width=3)
    D.text(s, xs[2] + cw, 5.5, gap, 0.3, ["read"], size=12, align="c", bold=True)
    # engine -> evaluation
    D.line(s, xs[2] + cw / 2, 2.3 + ch, xs[2] + cw / 2, 3.95, width=3, dash=True)
    D.line(s, xs[2] + cw / 2, 3.95, xs[3] + cw / 2, 3.95, width=3, dash=True)
    D.line(s, xs[3] + cw / 2, 3.95, xs[3] + cw / 2, 5.32, width=3, dash=True, arrow=True)
    D.text(s, xs[2] + cw / 2 + 0.1, 3.58, 2.6, 0.3, ["EvidenceRecord to score"], size=12, bold=True)
    D.text(s, 10.9, 5.35, 1.9, 1.2, ["Frontend: React + TypeScript; no backend"], size=12, color=SOFT, italic=True)

    # ============================================================ 6 TOPOLOGY
    s = D.slide("Network topology and testbed",
                "Six Linux network namespaces in a line: client, four routers, server, joined by veth pairs on /30 subnets. The prober and the "
                "application client both run in the client namespace so they read the same monotonic clock. Faults are injected on a router's "
                "egress link, so the first hop that can show the fault is the next one.")
    names = ["ns-client", "ns-r1", "ns-r2", "ns-r3", "ns-r4", "ns-server"]
    roles = ["client", "router", "router", "router", "router", "server"]
    fills = [LILAC, MINT100, MINT100, MINT100, MINT100, CYAN_T]
    addr = [["10.10.1.1"], ["10.10.1.2", "10.10.2.1"], ["10.10.2.2", "10.10.3.1"], ["10.10.3.2", "10.10.4.1"], ["10.10.4.2", "10.10.5.1"], ["10.10.5.2"]]
    nw, nh = 1.55, 1.05
    gap = (12.1 - 6 * nw) / 5
    ny = 2.45
    xsn = [0.6 + i * (nw + gap) for i in range(6)]
    for i in range(5):
        D.line(s, xsn[i] + nw, ny + nh / 2, xsn[i + 1], ny + nh / 2, width=5)
        D.text(s, xsn[i] + nw, ny + nh / 2 - 0.4, gap, 0.3, [f"s{i+1}"], size=12, mono=True, bold=True, align="c")
    for i, x in enumerate(xsn):
        D.card(s, x, ny, nw, nh, fill=fills[i], name=names[i])
        D.text(s, x + 0.05, ny + 0.1, nw - 0.1, nh - 0.2, [[(names[i], {"bold": True, "size": 17})], [(roles[i], {"size": 14, "color": SOFT})]],
               align="c", anchor="m")
        if i >= 1:
            D.pill(s, x + nw / 2 - 0.45, ny - 0.5, 0.9, 0.36, f"Hop {i}", size=14)
        else:
            D.text(s, x, ny - 0.45, nw, 0.3, ["path start"], size=13, color=SOFT, align="c")
        D.text(s, x - 0.05, ny + nh + 0.08, nw + 0.1, 0.6, [[(a, {})] for a in addr[i]], size=12, mono=True, color=SOFT, align="c")
    # lower annotations
    D.card(s, 0.6, 4.55, 3.9, 1.1, fill=WHITE, name="prober box")
    D.text(s, 0.75, 4.6, 3.6, 1.0, [[("Path prober + app client", {"bold": True})], [("co-located: one CLOCK_MONOTONIC", {"size": 13, "color": SOFT})]], size=16, anchor="m")
    D.line(s, xsn[0] + nw / 2, ny + nh + 0.7, xsn[0] + nw / 2, 4.55, width=2, dash=True)
    fx = (xsn[2] + nw + xsn[3]) / 2
    D.card(s, fx - 1.9, 4.55, 3.8, 1.1, fill=CORAL_T, name="fault box")
    D.text(s, fx - 1.8, 4.6, 3.6, 1.0, [[("Fault controller", {"bold": True})], [("tc netem/tbf on ns-r2 egress (example)", {"size": 13})]], size=16, anchor="m", align="c")
    D.line(s, fx, 4.55, fx, ny + nh / 2 + 0.07, color=DEGRADED, width=3, dash=True, arrow=True)
    D.card(s, 8.8, 4.55, 3.9, 1.1, fill=WHITE, name="server box")
    D.text(s, 8.95, 4.6, 3.6, 1.0, [[("Application server", {"bold": True})], [("WebSocket echo in ns-server", {"size": 13, "color": SOFT})]], size=16, anchor="m", align="r")
    D.line(s, xsn[5] + nw / 2, ny + nh + 0.7, xsn[5] + nw / 2, 4.55, width=2, dash=True)
    D.card(s, 0.6, 5.95, 12.1, 0.95, fill=YEL_T, dash=True, shadow=False, name="rule")
    D.text(s, 0.85, 5.98, 11.6, 0.9, [[("“Hop N onward” rule: ", {"bold": True}),
                                       ("a fault on the link leaving Hop N is first visible at Hop N+1, because Hop N's own reply is generated before that link. "
                                        "Verified here: 200/200 packets, hops 10.10.1.2 → 10.10.5.2 in order.", {})]], size=15, anchor="m")

    # ============================================================ 7 DATA COLLECTION
    s = D.slide("Data collection: two streams, one clock",
                "Parameters of both measurement streams. Both processes run in the client namespace and call the same monotonic clock function, "
                "so their timestamps can be compared directly with no clock synchronisation. Unanswered probes are recorded as not responded and never interpolated.")
    cw = 3.85
    cols3 = [("Path prober", CYAN_T, ["5 hops (TTL 1-5)", "2 probes per hop per round", "A round every 0.5 s", "0.3 s timeout; TCP-SYN (default) or UDP", "No answer = unknown, never interpolated"]),
             ("Application client", LILAC, ["Hand-written RFC 6455 WebSocket", "Ping at 20 Hz (feeds the detector)", "Interactions: Poisson, mean 2/s, 150 ms budget", "Missing reply after 2 s: RTT stays null", "RTT computed client-side only"]),
             ("One shared clock", YEL_T, ["Both in ns-client", "One CLOCK_MONOTONIC call", "Onsets directly comparable", "No cross-host synchronisation", "Bucket width 0.5 s = probe round"])]
    for i, (h1, fill, bl) in enumerate(cols3):
        x = 0.6 + i * (cw + 0.275)
        D.card(s, x, 1.45, cw, 3.55, fill=fill)
        D.text(s, x + 0.2, 1.52, cw - 0.4, 0.5, [h1], size=21, bold=True)
        D.text(s, x + 0.2, 2.1, cw - 0.4, 2.85, [[("•  ", {"bold": True}), (b, {})] for b in bl], size=16, space_after=5)
    D.text(s, 0.6, 5.1, 6.0, 0.3, ["hop.jsonl  (one line per probe)"], size=13, mono=True, bold=True, color=SOFT)
    D.code(s, 0.6, 5.42, 5.95, 1.45, ['{"round_id":0,"ttl":1,"hop_ip":"10.10.1.2",', ' "rtt_ms":0.050222,"responded":true,"probe_index":0,', ' "t_send_mono_ns":16726703720,', ' "t_recv_mono_ns":16726784215}'], size=12)
    D.text(s, 6.75, 5.1, 6.0, 0.3, ["app.jsonl  (one line per application sample)"], size=13, mono=True, bold=True, color=SOFT)
    D.code(s, 6.75, 5.42, 5.95, 1.45, ['{"seq":0,"kind":"ping",', ' "t_send_mono_ns":16793313189,', ' "t_recv_mono_ns":16793623395,', ' "rtt_ms":0.310206,"responded":true,"missed_budget":false}'], size=12)

    # ============================================================ 8 METHOD
    s = D.slide("Correlation methodology",
                "Left: the three operations. Bucket both streams into 500 millisecond buckets; threshold is median plus five MAD of the whole series and "
                "an onset needs three consecutive buckets above it; then align the application onset with hop onsets inside 1.5 seconds and require every deeper "
                "hop to agree. Right: the recorded demo run, recomputed with the project's own code. Hop 2 has one noisy bucket that the three-in-a-row rule ignores. "
                "Limitations: the baseline is global, so a fault that covers most of the window can hide itself.")
    items = [("1", "Bucket", "Mean RTT per 500 ms bucket for the app and for each hop. No answer = unknown."),
             ("2", "Onset", "T = median + 5 × MAD over the series. Onset = first of 3 consecutive buckets above T."),
             ("3", "Align", "|t_hop − t_app| ≤ 1.5 s, and every deeper hop must also show a concurrent onset.")]
    for i, (n, h1, t) in enumerate(items):
        y = 1.5 + i * 1.5
        D.card(s, 0.6, y, 5.6, 1.3, fill=WHITE)
        D.circle(s, 1.15, y + 0.65, 0.65, n, [UNKNOWN, SIGNAL, HEALTHY][i], size=20)
        D.text(s, 1.65, y + 0.06, 4.45, 0.4, [h1], size=19, bold=True)
        D.text(s, 1.65, y + 0.46, 4.45, 0.82, [t], size=15)
    D.card(s, 0.6, 6.0, 5.6, 0.9, fill=CORAL_T, dash=True, shadow=False, name="caveat")
    D.text(s, 0.75, 6.02, 5.35, 0.86, [[("Caveat: ", {"bold": True}), ("global baseline: a fault covering most of the window can yield “no onset”.", {})]], size=13.5, anchor="m")
    D.picture(s, "fig_demo_run_series.png", 6.55, 1.45, h=5.0, alt="Measured RTT series for application, Hop 2 and Hop 3 from demo-run-01", folder=FIGS)
    D.text(s, 6.55, 6.5, 6.2, 0.5, ["Measured: demo-run-01. App and Hop 3 onset at +25.0 s; Hop 2's single noisy bucket is ignored."], size=12, color=SOFT, italic=True)

    # ============================================================ 9 EVIDENCE
    s = D.slide("Evidence model and confidence",
                "Every claim carries exactly one of five tags. Confidence is a fixed heuristic: corroborating fraction times one minus unknown hops over five. "
                "Because persistence already requires all downstream hops to agree, the fraction is one whenever a range is reported, so the level depends only on "
                "how many hops were unknown. It is not a probability and it is not calibrated. Unknown is not healthy.")
    tags = [("fact", "measured", SIGNAL), ("inference", "consistent with", STRUCT), ("correlation", "timing only", HEALTHY),
            ("confidence", "heuristic", UNKNOWN), ("unknown", "not observed", WHITE)]
    cw, gap = 2.18, 0.3
    for i, (t, sub, fill) in enumerate(tags):
        x = 0.6 + i * (cw + gap)
        D.card(s, x, 1.5, cw, 1.35, fill=WHITE)
        D.pill(s, x + 0.2, 1.68, cw - 0.4, 0.4, t, fill=fill, size=15, mono=True, dash=(t == "unknown"), color=WHITE if t == "inference" else INK)
        D.text(s, x + 0.1, 2.2, cw - 0.2, 0.5, [sub], size=16, align="c", color=SOFT)
    D.card(s, 0.6, 3.15, 6.9, 3.7, fill=WHITE, name="formula card")
    D.text(s, 0.85, 3.22, 6.5, 0.45, ["Confidence: a fixed heuristic"], size=20, bold=True)
    D.text(s, 0.85, 3.8, 6.4, 0.9, [[("score = corroborating fraction × (1 − unknown hops / 5)", {"mono": True, "size": 15})]], size=15)
    D.table(s, 0.85, 4.55, 6.4, [2.4, 2.0, 2.0], [["unknown hops", "score", "level"], ["0", "1.0", "high"], ["1", "0.8", "high"], ["2", "0.6", "medium"], ["3", "0.4", "low"]],
            size=15, row_h=0.42,
            fills=[[None] * 3, [None, None, "E6FFC2"], [None, None, "E6FFC2"], [None, None, "FFF0B8"], [None, None, "FFE3DC"]])
    D.card(s, 7.8, 3.15, 4.9, 3.7, fill=YEL_T, dash=True, shadow=False, name="warning card")
    D.text(s, 8.05, 3.22, 4.5, 0.45, ["Read it correctly"], size=20, bold=True)
    D.text(s, 8.05, 3.8, 4.45, 3.0, [[("•  ", {"bold": True}), (t, {})] for t in
           ["Not a probability; not calibrated", "Ignores signal strength", "Unknown ≠ healthy", "“No onset” ≠ “no fault”", "Only high confidence prints “…at or after Hop N”"]],
           size=17, space_after=7)

    # ============================================================ 10 STACK
    s = D.slide("Implementation and technology stack",
                "What is actually used. Python code is standard library only, with no third-party packages. The testbed runs on WSL2 Ubuntu 24.04. The dashboard is React, "
                "TypeScript, Vite, Tailwind, Zod and Motion. Required versus optional tools are separated in the report.")
    stats = [("0", "third-party\nPython packages"), ("6", "network\nnamespaces"), ("5", "hops\nmeasured"), ("2", "OS identities\nisolate labels")]
    for i, (big, lab) in enumerate(stats):
        x = 0.6 + i * 3.08
        D.card(s, x, 1.45, 2.85, 1.45, fill=[CYAN_T, LILAC, CORAL_T, YEL_T][i])
        D.text(s, x + 0.1, 1.5, 1.1, 1.35, [big], size=44, font=HEAD, anchor="m", align="c")
        D.text(s, x + 1.15, 1.5, 1.65, 1.35, lab.split("\n"), size=15.5, anchor="m")
    cols3 = [("Measure and control", ["Python 3.12.3, stdlib only", "Raw sockets (TCP-SYN / UDP)", "RFC 6455 WebSocket, hand-written", "iproute2 6.1: ip netns, tc netem/tbf", "SQLite ground-truth store"]),
             ("Analyse and present", ["synchronize → onset → align → evidence", "Locked result sentences", "path_bundle.json export", "Metrics: Wilson CI, Brier score", "unittest: 117 tests"]),
             ("Dashboard", ["React 19 · TypeScript (strict)", "Vite 8 · Tailwind 4 · Motion", "Zod schema validation", "Vitest: 310 tests; axe, ESLint a11y", "Node ≥ 22.12; no CDN at runtime"])]
    for i, (h1, bl) in enumerate(cols3):
        x = 0.6 + i * 4.1
        D.card(s, x, 3.2, 3.9, 3.0, fill=WHITE)
        D.text(s, x + 0.2, 3.27, 3.5, 0.45, [h1], size=19, bold=True)
        D.text(s, x + 0.2, 3.8, 3.55, 2.4, [[("•  ", {"bold": True}), (b, {})] for b in bl], size=15.5, space_after=4)
    D.text(s, 0.6, 6.4, 12.1, 0.5, [[("Environment: ", {"bold": True}), ("WSL2 Ubuntu 24.04 (kernel 6.18.33.2) for the testbed; Windows 11 + Node 24.14.1 for the dashboard checks. Windows-native Python is unsupported.", {})]], size=14, color=SOFT)

    # ============================================================ 11 DASHBOARD 1
    s = D.slide("Dashboard: from symptom to hop",
                "Genuine screenshots of the production build, demo-run-01. Left: the application's own RTT, the symptom, with the stored before and after means and the "
                "detected onset. Right: the path with a stored status per hop, and the hop inspector for Hop 3. The moving packets are illustrative only; the data contain no "
                "packet-level observations. The dashboard displays stored results and never recomputes them.")
    D.picture(s, "02-application-rtt-chart.png", 1.45, 1.4, h=4.4, alt="Dashboard: application RTT chart for demo-run-01")
    D.picture(s, "03-path-and-hop-inspector.png", 7.26, 1.4, h=4.4, alt="Dashboard: path with hop statuses and hop inspector for Hop 3")
    D.circle(s, 1.7, 6.3, 0.42, "1", SIGNAL, size=13)
    D.circle(s, 7.5, 6.3, 0.42, "2", UNKNOWN, size=13)
    D.text(s, 2.05, 6.0, 4.9, 0.9, [[("The symptom. ", {"bold": True}), ("App RTT 0.633 ms → 90.4 ms; onset at +25.0 s (700 pings).", {})]], size=14)
    D.text(s, 7.85, 6.0, 4.9, 0.9, [[("The path. ", {"bold": True}), ("Hops 3-5 corroborated; Hop 3 inspected: 0.015 → 90.3 ms. Moving packets are ", {}), ("illustrative only", {"bold": True}), (".", {})]], size=14)

    # ============================================================ 12 DASHBOARD 2
    s = D.slide("Dashboard: unknowns, evidence, guided mode",
                "Left: a held-out trial where Hops 1 and 2 returned too few responses; they are shown as unknown, not healthy. Middle: evidence kept apart by kind, "
                "with the confidence panel stating it is a heuristic and not a probability. Right: the optional guided explanation, off by default, built only from the selected bundle.")
    from PIL import Image as _Im
    shots = [("04-unknown-hops-trial.png", "Dashboard: trial with unknown hops 1 and 2", "Unknown ≠ healthy", "p4-held-12: Hops 1-2 unknown (hatched yellow); partially localized, medium confidence."),
             ("07-evidence-by-kind.png", "Dashboard: evidence grouped by kind and confidence panel", "Evidence, kept apart", "Facts, inference, correlation, unknown; confidence panel: “not a probability”."),
             ("08-guided-explanation-overview.png", "Dashboard: guided explanation overview", "Guided explanation", "Off by default; five-step walkthrough from stored data only.")]
    hh, xx = 3.0, 0.6
    avail = 12.1 - 2 * 0.3
    widths = []
    for f, *_ in shots:
        iw, ih = _Im.open(os.path.join(SHOTS, f)).size
        widths.append(iw / ih * hh)
    scale = min(1.0, avail / sum(widths))
    for (f, alt, h1, t), w in zip(shots, widths):
        w *= scale
        D.picture(s, f, xx, 1.4, w=w, alt=alt)
        D.text(s, xx, 4.55, w, 1.5, [[(h1, {"bold": True, "size": 17})], [(t, {"size": 14.5})]], size=14.5, space_after=3)
        xx += w + 0.3
    D.card(s, 0.6, 6.15, 12.1, 0.75, fill=CYAN_T, shadow=False, name="dash note")
    D.text(s, 0.85, 6.17, 11.6, 0.7, [[("Screens captured from the current build. ", {"bold": True}), ("The page only formats stored values; statuses, onsets and confidence come from the bundle.", {})]], size=14.5, anchor="m")

    # ============================================================ 13 TESTING
    s = D.slide("Testing and verification (run 27 Oct 2026)",
                "These checks were actually run for this submission. Unit tests are separate from integration or browser checks and from the experimental evaluation. "
                "The right-hand card lists what was deliberately not run: the live demo overwrites a recorded trial, the evaluation scripts overwrite result files, and offline, "
                "cross-browser and clean-clone tests have not been done.")
    stats = [("117", "Python unit tests\nall OK (WSL2)", CYAN_T), ("310", "frontend tests\n14 files, all pass", LILAC), ("40", "layout checks\n10 widths × 4 trials", HEALTHY), ("200/200", "packets, 0% loss\n5 hops in order", UNKNOWN)]
    for i, (big, lab, fill) in enumerate(stats):
        x = 0.6 + i * 3.08
        D.card(s, x, 1.45, 2.85, 1.7, fill=fill)
        D.text(s, x + 0.1, 1.5, 2.65, 0.85, [big], size=34 if len(big) > 4 else 42, font=HEAD, align="c", anchor="m")
        D.text(s, x + 0.1, 2.35, 2.65, 0.75, lab.split("\n"), size=14, align="c")
    D.card(s, 0.6, 3.45, 6.9, 3.4, fill=WHITE, name="also pass")
    D.text(s, 0.85, 3.52, 6.4, 0.45, ["Also verified"], size=19, bold=True)
    D.text(s, 0.85, 4.0, 6.5, 2.8, [[("✓  ", {"bold": True, "color": "2E7D32"}), (t, {})] for t in
           ["TypeScript strict typecheck; ESLint, 0 warnings", "Production build + distribution-safety check", "4 trials switched in a real browser; packet pause; guided walkthrough",
            "Repository hygiene check; live testbed build + teardown", "Report and deck numbers asserted against result files"]], size=16.5, space_after=7)
    D.card(s, 7.8, 3.45, 4.9, 3.4, fill=YEL_T, dash=True, shadow=False, name="not run")
    D.text(s, 8.05, 3.52, 4.5, 0.45, ["Not run (disclosed)"], size=19, bold=True)
    D.text(s, 8.05, 4.0, 4.45, 2.8, [[("•  ", {"bold": True}), (t, {})] for t in
           ["Full live demo (overwrites a recorded trial)", "Historical evaluation scripts (overwrite results)", "Offline and cross-browser tests", "Clean-clone test"]],
           size=16.5, space_after=8)

    # ============================================================ 14 RESULTS
    s = D.slide("Results: Phase 4 failed its target",
                "State this plainly. Phase 4 was the only independent held-out run and it missed the false-attribution target: 3 of 12, 25 percent, against at most 10 percent. "
                "Phase 4b scored 1 of 12 on the same held-out set after its traces had been inspected during design, so it is a historical record, not independent validation. "
                "The table compares both. All samples are tiny: 12 trials, 8 faults and 4 clean controls.")
    D.card(s, 0.6, 1.4, 5.95, 2.35, fill=CORAL_T, name="phase4")
    D.text(s, 0.85, 1.45, 5.5, 0.4, ["Phase 4 · mad · INDEPENDENT held-out run"], size=14.5, bold=True, mono=True)
    D.text(s, 0.85, 1.85, 3.1, 1.2, ["3/12"], size=60, font=HEAD, anchor="m")
    D.text(s, 3.85, 1.95, 2.6, 1.1, [[("= 25%", {"size": 30, "bold": True})], [("false attribution", {"size": 14})]], size=14, anchor="m")
    D.pill(s, 0.85, 3.12, 5.4, 0.45, "✕  TARGET MISSED: at most 10%", fill=DEGRADED, size=15, color=INK)
    D.card(s, 6.8, 1.4, 5.9, 2.35, fill=YEL_T, dash=True, shadow=False, name="phase4b")
    D.text(s, 7.05, 1.45, 5.5, 0.4, ["Phase 4b · mad_sustained · EXPOSED set"], size=14.5, bold=True, mono=True)
    D.text(s, 7.05, 1.85, 3.1, 1.2, ["1/12"], size=60, font=HEAD, anchor="m")
    D.text(s, 10.05, 1.95, 2.6, 1.1, [[("= 8.3%", {"size": 30, "bold": True})], [("false attribution", {"size": 14})]], size=14, anchor="m")
    D.pill(s, 7.05, 3.12, 5.4, 0.45, "⚠  same trials, inspected during design", fill=UNKNOWN, size=15, dash=True)
    rows = [["Metric (12 trials)", "Phase 4 · independent", "Phase 4b · exposed"],
            ["Within ±1 hop (localized trials)", "3/3", "4/4  (95% CI 51–100%)"],
            ["Exact first hop (localized trials)", "2/3", "3/4"],
            ["Clean-control false positives", "2/4", "0/4"],
            ["Recall of injected faults", "3/8 = 37.5%", "4/8 = 50%"],
            ["Brier score (confidence check)", "0.432  (n = 5)", "0.040  (n = 4)"]]
    fills = [[None] * 3] + [[None, None, None] for _ in range(5)]
    fills[3][1] = "FFD9D0"; fills[4][1] = "FFF0B8"
    D.table(s, 0.6, 3.95, 12.1, [4.9, 3.5, 3.7], rows, size=16, row_h=0.45, fills=fills)
    D.text(s, 0.6, 6.7, 12.1, 0.3, ["Meeting a numeric threshold on an exposed set does not validate the current detector."], size=12.5, italic=True, color=SOFT)

    # ============================================================ 15 PER-TRIAL
    s = D.slide("What the held-out trials show",
                "Per-trial outcomes on the held-out set. Latency and jitter were localized when a claim was made. Loss and bandwidth: no claim in any of the four trials, under either detector, "
                "and none on the Phase 4b tuning set. The remaining false attribution, trial 12, occurred where the first affected hop was itself unknown. Four clean controls cannot separate zero "
                "from nearly fifty percent false positives.")
    tr = [("01 clean", "clean", "FA", "–"), ("02 clean", "clean", "–", "–"), ("03 clean", "clean", "FA", "–"),
          ("04 latency · hop 2", "latency", "✓", "✓"), ("05 jitter · hop 3", "jitter", "✓", "✓"), ("06 loss · hop 4", "loss", "miss", "miss"),
          ("07 bandwidth · hop 1", "bandwidth", "miss", "miss"), ("08 jitter · hop 2, transient", "jitter", "miss", "✓"),
          ("09 bandwidth · hop 3, transient", "bandwidth", "miss", "miss"), ("10 loss · hop 1", "loss", "miss", "miss"),
          ("11 clean · Hop 1 silent", "clean", "–", "–"), ("12 jitter · Hops 1-2 silent", "jitter", "FA", "FA")]
    rows = [["Held-out trial", "Phase 4", "Phase 4b"]] + [[a, c, d] for a, _, c, d in tr]
    cmap = {"FA": "FFD9D0", "miss": "FFF0B8", "✓": "E6FFC2", "–": None}
    fills = [[None] * 3] + [[None, cmap[c], cmap[d]] for _, _, c, d in tr]
    D.table(s, 0.6, 1.4, 6.9, [3.9, 1.5, 1.5], rows, size=14, row_h=0.4, fills=fills)
    D.text(s, 0.6, 6.7, 6.9, 0.3, ["✓ correct claim   FA false attribution   miss no claim   – no claim, correctly"], size=12, color=SOFT)
    take = [(HEALTHY, "Latency and jitter localized", "when a claim was made (one jitter claim was a false attribution)"),
            (DEGRADED, "Loss and bandwidth: 0 of 4", "held-out; also 0 of 4 on the Phase 4b tuning set"),
            (UNKNOWN, "Small samples", "one trial moves a rate by 8.3 points; 4/4 has a 95% interval of 51–100%")]
    for i, (fill, h1, t) in enumerate(take):
        y = 1.4 + i * 1.8
        D.card(s, 7.8, y, 4.9, 1.6, fill=WHITE)
        D.circle(s, 8.3, y + 0.8, 0.55, str(i + 1), fill, size=16)
        D.text(s, 8.7, y + 0.1, 3.9, 1.45, [[(h1, {"bold": True, "size": 17})], [(t, {"size": 14})]], size=14, anchor="m", space_after=2)

    # ============================================================ 16 LIMITATIONS
    s = D.slide("Limitations and threats to validity",
                "Six limitations, each supported by a file in the repository. The most important: no independent validation of the current detector. Also loss and bandwidth are not localized, "
                "confidence is an uncalibrated heuristic, unknown hops cannot be assumed healthy, timing correlation is not causation, and the samples and parameters are small and fixed.")
    lim = [("≠", DEGRADED, "No independent validation", "Phase 4 missed its target; Phase 4b ran on an exposed set; the sealed evaluation (spec Phase 5) has not been run"),
           ("0/4", UNKNOWN, "Loss and bandwidth", "Not localized (held-out and tuning). Bandwidth observability with current probe load is unproven"),
           ("≈", SIGNAL, "Confidence is a heuristic", "Fixed formula, uncalibrated, not a probability; it falls only when hops are unknown"),
           ("?", HEALTHY, "Unknown ≠ healthy", "No onset can mean healthy, a long fault, or too little data; there is no “insufficient data” state"),
           ("→", STRUCT, "Correlation ≠ causation", "Timing evidence only. The 1.5 s window and 500 ms resolution are fixed and unvalidated"),
           ("12", LILAC, "Small, controlled setting", "12 + 12 + 10 trials vs a 120-trial spec; one linear path; ICMP rate limit switched off")]
    cw, ch = 3.9, 2.55
    for i, (sym, fill, h1, t) in enumerate(lim):
        x = 0.6 + (i % 3) * (cw + 0.2)
        y = 1.4 + (i // 3) * (ch + 0.25)
        D.card(s, x, y, cw, ch, fill=WHITE)
        D.circle(s, x + 0.65, y + 0.65, 0.85, sym, fill, size=18 if len(sym) < 3 else 15, color=WHITE if fill == STRUCT else INK)
        D.text(s, x + 1.2, y + 0.2, cw - 1.35, 0.9, [h1], size=17, bold=True, anchor="m")
        D.text(s, x + 0.25, y + 1.2, cw - 0.5, 1.3, [t], size=15)

    # ============================================================ 17 CONCLUSION
    s = D.slide("Conclusion and future work",
                "The engineering works end to end and every check run for this submission passes. The science is less settled and the project says so: the independent evaluation missed its "
                "target, loss and bandwidth are not localized, confidence is a heuristic. Future work is recorded but not yet authorized.")
    D.card(s, 0.6, 1.45, 6.0, 5.4, fill=WHITE, name="conclusion")
    D.text(s, 0.85, 1.52, 5.5, 0.45, ["What HOPTRACE shows"], size=21, bold=True)
    D.text(s, 0.85, 2.1, 5.55, 4.7, [[("•  ", {"bold": True}), (t, {})] for t in
           ["App RTT and per-hop RTT can be compared on one shared clock to estimate where degradation appears to begin",
            "It runs end to end: testbed, prober, WebSocket app, isolated fault controller, engine, dashboard",
            "It is honest about failure: the independent run missed its target and the interface says so",
            "It claims timing evidence, never a cause"]], size=18, space_after=12)
    D.card(s, 6.9, 1.45, 5.8, 5.4, fill=CYAN_T, name="future")
    D.text(s, 7.15, 1.52, 5.3, 0.45, ["Future work (recorded, not yet authorized)"], size=19, bold=True)
    D.text(s, 7.15, 2.1, 5.35, 4.7, [[(f"{i+1}  ", {"bold": True, "mono": True}), (t, {})] for i, t in enumerate(
        ["Sealed evaluation set + calibration check (spec Phase 5)", "Is bandwidth observable with controlled background traffic?",
         "An explicit “insufficient data” result state", "Validate the 1.5 s window and global-baseline assumption",
         "Offline, cross-browser and clean-clone tests"])], size=18, space_after=12)

    # ============================================================ 18 REFERENCES
    s = D.slide("References and questions",
                "Close with the sources the implementation relies on and the deliverables: the report PDF, this deck, and the GitHub repository once publication is approved.")
    refs = ["RFC 6455: The WebSocket Protocol (Fette, Melnikov), 2011", "RFC 792: Internet Control Message Protocol (Postel), 1981",
            "RFC 9293: Transmission Control Protocol; RFC 1071: Internet checksum",
            "Linux tc-netem(8), tc-tbf(8), ip-netns(8), veth(4), clock_gettime(2)",
            "Hampel 1974 (MAD); Page 1954 (CUSUM); Adams & MacKay 2007 (BOCPD); Killick et al. 2012 (PELT)",
            "Wilson 1927 (score interval); Brier 1950 (Brier score)",
            "Project: hoptrace-phase-0-5-spec.html, docs/EVIDENCE_LEDGER.md"]
    D.card(s, 0.6, 1.45, 7.6, 5.4, fill=WHITE, name="refs")
    D.text(s, 0.85, 1.52, 7.1, 0.45, ["References (full list in the report)"], size=19, bold=True)
    D.text(s, 0.85, 2.05, 7.15, 4.75, [[("•  ", {"bold": True}), (r, {})] for r in refs], size=16, space_after=8)
    D.card(s, 8.5, 1.45, 4.2, 2.6, fill=LILAC, name="deliverables")
    D.text(s, 8.7, 1.52, 3.85, 0.45, ["Deliverables"], size=19, bold=True)
    D.text(s, 8.7, 2.05, 3.85, 2.0, [[("Report: ", {"bold": True}), ("PDF, Assignment 8", {})], [("Deck: ", {"bold": True}), ("this file, Assignment 9", {})],
                                      [("Code: ", {"bold": True}), (f"GitHub repo {slot}_Team{team}, Assignment 10", {})],
                                      [(f"https://github.com/KrishUsesGit/{slot}_Team{team}", {"mono": True, "size": 11, "link": f"https://github.com/KrishUsesGit/{slot}_Team{team}"})]], size=16, space_after=8)
    D.card(s, 8.5, 4.3, 4.2, 2.55, fill=INK, outline=INK, shadow=False, name="questions")
    D.text(s, 8.7, 4.4, 3.85, 2.3, [[("Questions?", {"size": 34, "font": HEAD, "color": WHITE})],
                                      [("Krish Chirag Shah · 24BCE0624", {"size": 14, "color": SIGNAL, "mono": True})]], size=14, anchor="m", space_after=8)

    # ---------------------------------------------------------------- save
    name = f"{slot}_Team{team}_PPT.pptx"
    out = os.path.join(ROOT, "submission", name)
    D.prs.core_properties.title = "HOPTRACE: Network Path Fault Localization through Application-to-Hop Correlation"
    D.prs.core_properties.author = "Krish Chirag Shah"
    D.prs.core_properties.last_modified_by = "Krish Chirag Shah"
    import datetime as _dt
    _now = _dt.datetime.now(_dt.timezone.utc).replace(tzinfo=None, microsecond=0)
    D.prs.core_properties.created = _now
    D.prs.core_properties.modified = _now
    D.prs.core_properties.revision = 1
    D.prs.core_properties.comments = ""
    D.prs.core_properties.subject = "BCSE308P Computer Networks Lab, Assignment 9"
    D.prs.save(out)
    print(f"wrote {out} ({D.n} slides)")
    return out


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--slot", default="F2")
    ap.add_argument("--team", default="24")
    a = ap.parse_args()
    build(a.slot, a.team)
